"""Pruebas de aceptacion del checklist: briefing (E1-E17), normativa, catalogo, buscar, doc y verificar.

Corpus sintetico de dos areas (Seguridad y Calidad) con DES, PRD, PRO y POL, y varios POL-GEN: uno que
menciona al area, uno que menciona a un puesto, uno cuyo alcance es toda la empresa y uno sin relacion.
"""

import csv
import io
import os
import re
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from tests import ayuda  # noqa: F401
from corpus_org.consulta import cli as consulta_cli
from corpus_org.conversion import cli as conversion_cli

try:
    import docx  # noqa: F401
    import lxml  # noqa: F401
    import reportlab  # noqa: F401
    from tests.fixtures_corpus import desc, proc
    DEPENDENCIAS = True
except ImportError:
    DEPENDENCIAS = False

PASO_GEN = [("Revisar el documento", "Personal administrativo")]
PASOS = [("Realizar la ronda", "Jefe de Seguridad"), ("Aprobar el informe de la ronda", "Jefe de Seguridad"),
         ("Registrar novedades", "Vigilante")]
TITULOS_BRIEFING = ["Puestos (DES)", "Procedimientos (PRD)", "Protocolos (PRO)", "Políticas (POL)",
                    "Jerarquía de procesos", "Relaciones", "Participación — Resumen por puesto",
                    "Procedimientos con un solo puesto participante",
                    "Mismo puesto con roles Ejecuta y Aprueba en un procedimiento", "Estructura",
                    "Indicadores relacionados", "Gobierno documental (resumen)", "Normativa citada (resumen)",
                    "Perfil y funciones (resumen por puesto)", "No documentado"]


def proc_con_responsabilidades(path, header, objetivo, pasos, responsabilidades):
    """Como ``proc`` pero con el texto de la seccion Responsabilidades indicado."""
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_heading("Objetivo", 1)
    d.add_paragraph(objetivo)
    d.add_heading("Alcance", 1)
    d.add_paragraph("Aplica a las instalaciones de la empresa.")
    d.add_heading("Responsabilidades", 1)
    d.add_paragraph(responsabilidades)
    d.add_heading("Desarrollo del proceso", 1)
    t = d.add_table(rows=len(pasos) + 1, cols=3)
    for j, h in enumerate(["No.", "Actividad", "Responsable"]):
        t.cell(0, j).text = h
    for i, (a, r) in enumerate(pasos, 1):
        t.cell(i, 0).text = str(i)
        t.cell(i, 1).text = a
        t.cell(i, 2).text = r
    d.save(path)


def _consultar(copilot, *argv):
    flujo = io.StringIO()
    with mock.patch.dict(os.environ, {"CORPUS_DATOS": str(copilot)}):
        codigo = consulta_cli.main(list(argv), uso="uso", flujo=flujo)
    return codigo, flujo.getvalue()


@unittest.skipUnless(DEPENDENCIAS, "faltan dependencias de conversion o de las pruebas")
class BriefingAceptacionTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        base = Path(cls.tmp.name)
        ent, ctx = base / "entrada", base / "contexto"
        ctx.mkdir(parents=True)

        def ruta(*partes):
            p = ent.joinpath(*partes)
            p.parent.mkdir(parents=True, exist_ok=True)
            return p
        desc(ruta("Descriptivos", "Seguridad", "Jefe de Seguridad.docx"), "Jefe de Seguridad", "Gerente General",
             "Proteger las instalaciones.", ["Coordinar las rondas.", "Aplicar la política de accesos."], "Código: DP-SEG-01")
        desc(ruta("Descriptivos", "Seguridad", "Vigilante.docx"), "Vigilante", "Jefe de Seguridad",
             "Vigilar los accesos.", ["Registrar las visitas."], "Código: DP-SEG-02")
        desc(ruta("Descriptivos", "Calidad", "Auditor de Calidad.docx"), "Auditor de Calidad", "Gerente General",
             "Auditar procesos.", ["Elaborar informes."], "Código: DP-CAL-01")
        proc_con_responsabilidades(ruta("Procedimientos", "Seguridad", "Rondas.docx"),
                                   "Código: PR-SEG-001 Versión: 1", "Recorrer las instalaciones.", PASOS,
                                   "El Jefe de Seguridad aprueba el informe de la ronda.")
        proc(ruta("Procedimientos", "Calidad", "Auditorias.docx"), "Código: PR-CAL-001",
             "Auditar los procesos.", [("Auditar", "Auditor de Calidad")], [])
        proc(ruta("Protocolos", "Seguridad", "Protocolo de emergencias.docx"), "Código: PT-SEG-001 Versión: 2",
             "Actuar ante emergencias.", PASOS, [])
        proc(ruta("Políticas", "Seguridad", "Política de accesos.docx"), "Código: PL-SEG-001",
             "Controlar los accesos.", PASOS, [])
        # POL-GEN: menciona al area / a un puesto / alcance toda la empresa / sin relacion
        proc(ruta("Políticas", "General", "Política de rondas.docx"), "Código: PL-GEN-001 Área propietaria: Calidad",
             "Regular las rondas de Seguridad en las sedes.", PASO_GEN, [])
        proc(ruta("Políticas", "General", "Política de excepciones.docx"), "Código: PL-GEN-002",
             "El Jefe de Seguridad aprueba las excepciones.", PASO_GEN, [])
        proc(ruta("Políticas", "General", "Política global.docx"), "Código: PL-GEN-003",
             "Establecer reglas comunes.", PASO_GEN, ["Aplica a toda la empresa."])
        proc(ruta("Políticas", "General", "Política de estacionamiento.docx"), "Código: PL-GEN-004",
             "Regular el uso del estacionamiento. Asegurar la calidad de vida.", PASO_GEN, [])
        proc(ruta("Protocolos", "General", "Protocolo de visitas.docx"), "Código: PT-GEN-001",
             "Recibir visitantes en las sedes.", PASO_GEN, [])
        (ctx / "kpis.csv").write_text("id;gestion;indicador;meta;medicion;terminos_busqueda\n"
                                      "K-01;Seguridad;Rondas realizadas;30;28;ronda|rondas\n"
                                      "K-02;Seguridad;Incidentes;0;;incidente\n", encoding="utf-8")
        with redirect_stdout(io.StringIO()) as out:
            try:
                conversion_cli.main(["construir", "--entrada", str(ent), "--contexto", str(ctx),
                                     "--sin-verificacion-cruzada", "--muestra-revision", "0"])
                cls.codigo = 0
            except SystemExit as e:
                cls.codigo = e.code
        cls.salida = out.getvalue()
        cls.copilot = ent / "salida_corpus" / "copilot"
        with open(cls.copilot / ayuda.INDICE / "indice_corpus.csv", encoding="utf-8-sig") as f:
            cls.codigo_de = {r["archivo_origen"]: r["codigo"] for r in csv.DictReader(f, delimiter=";")}
        cls.mapa = {}
        for p in (cls.copilot / "fuente-documentos").rglob("*.md"):
            cls.mapa[p.name] = p.read_text(encoding="utf-8")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_g1_verificar(self):
        codigo, out = _consultar(self.copilot, "verificar")
        self.assertEqual(codigo, 0, out)
        self.assertIn("Resultado: sin problemas", out)
        self.assertIn("Peso y archivos · fuente-contexto:", out)
        self.assertIn("Peso y archivos · fuente-documentos:", out)
        self.assertIn("corpus completo: sí", out)

    def test_g2_catalogo(self):
        _, out = _consultar(self.copilot, "catalogo")
        cab = next(l for l in out.splitlines() if l.startswith("| Código | Área |"))
        for t in ("DES", "PRD", "PRO", "POL"):
            self.assertIn(f"| {t} ", cab)
        self.assertRegex(out, r"\| SEG \| Seguridad \| 2 \| 1 \| 1 \| 1 \|")
        self.assertRegex(out, r"\| GEN \| General \| 0 \| 0 \| 1 \| 4 \|")      # los GEN, en su propia fila
        self.assertIn("Cobertura:", out)

    def test_g3_briefing_tablas_en_orden(self):
        codigo, out = _consultar(self.copilot, "briefing", "Seguridad")
        self.assertEqual(codigo, 0, out)
        lineas = out.splitlines()
        self.assertEqual(lineas[0], "Partes: 1")
        self.assertEqual(lineas[1], "Cobertura: DES 2/2 · PRD 1/1 · PRO 1/1 · POL 1/1 · GEN pertinentes: 3")
        posiciones = [out.index(f"\n## {t}\n") for t in TITULOS_BRIEFING]
        self.assertEqual(posiciones, sorted(posiciones), "las tablas no están en el orden del checklist")
        self.assertIn("N.º de puestos que participan", out)
        self.assertIn("| Verbo |", out)
        self.assertIn("Nivel jerárquico", out)

    def test_g4_gen_pertinentes_con_motivo(self):
        _, out = _consultar(self.copilot, "briefing", "Seguridad")
        politicas = out[out.index("## Políticas (POL)"):out.index("## Jerarquía de procesos")]
        gen = lambda archivo: self.codigo_de[archivo]                                                    # noqa: E731
        fila = lambda codigo: next(l for l in politicas.splitlines() if l.startswith(f"| {codigo} "))   # noqa: E731
        rondas, excepciones = gen("Política de rondas.docx"), gen("Política de excepciones.docx")
        self.assertEqual(fila(rondas).count("Menciona al área"), 1)
        self.assertIn("Menciona al puesto Jefe de Seguridad", fila(excepciones))
        self.assertIn("Alcance: toda la empresa", fila(gen("Política global.docx")))
        self.assertIn("Del área", fila(self.codigo_de["Política de accesos.docx"]))
        # sin relacion con el area: no entra por afinidad de tema
        self.assertNotIn(gen("Política de estacionamiento.docx"), out)
        self.assertNotIn(gen("Protocolo de visitas.docx"), out)
        self.assertIn("Calidad", fila(rondas))                    # el area propietaria queda aparte
        _, calidad = _consultar(self.copilot, "briefing", "Calidad")
        self.assertIn("Alcance: toda la empresa", calidad)
        self.assertNotIn(rondas, calidad.split("## Jerarquía de procesos")[0])

    def test_g3_resumenes_del_area(self):
        _, out = _consultar(self.copilot, "briefing", "Seguridad")
        # E12 estructura: el vigilante reporta al jefe de seguridad (nivel 2); el jefe reporta fuera (nivel 1)
        estructura = out[out.index("## Estructura"):out.index("## Indicadores relacionados")]
        self.assertRegex(estructura, r"\| DES-SEG-01 \| Jefe de Seguridad \| Gerente General \| — \| 1 \| DES-SEG-02 \| 1 \|")
        self.assertRegex(estructura, r"\| DES-SEG-02 \| Vigilante \| Jefe de Seguridad \| DES-SEG-01 \| 0 \| — \| 2 \|")
        # E13 indicadores: K-02 sin medicion; E14 gobierno; E15 normativa; E16 perfil
        indicadores = out[out.index("## Indicadores relacionados"):out.index("## Gobierno documental")]
        self.assertIn("K-01", indicadores)
        self.assertIn("Indicadores sin medición", indicadores)
        self.assertIn("Procedimientos del área sin ningún indicador relacionado", indicadores)
        gobierno = out[out.index("## Gobierno documental"):out.index("## Normativa citada")]
        self.assertRegex(gobierno, r"\| Versión \| \d+ de 5 \|")
        self.assertIn("| Tipo de referencia | N.º de citas | Documentos que la citan |", out)
        self.assertIn("Educación requerida", out)
        # E10 / E11: el Jefe de Seguridad ejecuta y aprueba en las rondas
        ejecuta = out[out.index("## Mismo puesto con roles Ejecuta y Aprueba"):out.index("## Estructura")]
        self.assertIn("DES-SEG-01", ejecuta)
        self.assertIn("Realizar la ronda", ejecuta)
        self.assertIn("aprueba el informe de la ronda", ejecuta)

    def test_normativa(self):
        _, out = _consultar(self.copilot, "normativa", "--area", "Seguridad")
        self.assertTrue(out.startswith("NORMATIVA CITADA"))
        self.assertIn("| Código | Documento | Tipo de referencia | Cita literal | Sección |", out)
        self.assertRegex(out, r"## Política \(\d+ citas")

    def test_g5_buscar_por_tipo(self):
        _, out = _consultar(self.copilot, "buscar", "emergencias|ronda|accesos", "--tipo", "PRO")
        codigos = set(re.findall(r"\b(?:DES|PRD|PRO|POL)-[A-Z]{3}-\d{2}\b", out))
        self.assertEqual({c.split("-")[0] for c in codigos}, {"PRO"}, out)
        _, gen = _consultar(self.copilot, "buscar", "sedes", "--tipo", "GEN")
        self.assertEqual({c.split("-")[1] for c in re.findall(r"\b(?:DES|PRD|PRO|POL)-[A-Z]{3}-\d{2}\b", gen)}, {"GEN"})

    def test_g6_doc(self):
        _, out = _consultar(self.copilot, "doc", "DES-SEG-01")
        self.assertIn("## [DES-SEG-01] Jefe de Seguridad", out)
        self.assertIn("Proteger las instalaciones.", out)

    def test_a1_titulos_y_sin_formato_anterior(self):
        titulo = re.compile(r"^## \[[A-Z]{3}-[A-Z]{3}-\d{2,3}\] ", re.M)
        viejo = re.compile(r"(?<![A-Za-z0-9])[DPC]-[A-Z0-9]{3}-\d{2,3}\b")
        for nombre, texto in self.mapa.items():
            for l in texto.splitlines():
                if l.startswith("## "):
                    self.assertRegex(l + " ", titulo, (nombre, l))
            self.assertIsNone(viejo.search(texto), nombre)
        self.assertIn("POL_GEN.md", self.mapa)
        self.assertIn("PRO_GEN.md", self.mapa)

    def test_citar_confirma_que_el_documento_existe_y_respalda(self):
        codigo, out = _consultar(self.copilot, "citar", "DES-SEG-01", "--texto", "Proteger las instalaciones")
        self.assertEqual(codigo, 0, out)
        self.assertIn("EXISTE · DES-SEG-01", out)
        self.assertIn("RESPALDADO (literal) · DES-SEG-01", out)
        self.assertIn("Específico del área", out)
        self.assertIn("entrada/Descriptivos/Seguridad/Jefe de Seguridad.docx", out)    # donde esta el original

    def test_citar_avisa_si_no_hay_respaldo_o_no_existe(self):
        codigo, out = _consultar(self.copilot, "citar", "DES-SEG-01", "--texto", "Administrar la nomina de la empresa")
        self.assertEqual(codigo, 3, out)
        self.assertIn("SIN RESPALDO · DES-SEG-01", out)
        codigo, out = _consultar(self.copilot, "citar", "DES-SEG-99")
        self.assertEqual(codigo, 2, out)
        self.assertIn("NO EXISTE · DES-SEG-99", out)

    def test_citar_un_documento_general_dice_que_no_es_especifico(self):
        codigo, out = _consultar(self.copilot, "citar", self.codigo_de["Política global.docx"])
        self.assertEqual(codigo, 0, out)
        self.assertIn("GENERAL (GEN): contexto o norma de toda la empresa", out)

    def test_el_indice_guarda_donde_esta_el_original(self):
        with open(self.copilot / ayuda.INDICE / "indice_corpus.csv", encoding="utf-8-sig") as f:
            filas = list(csv.DictReader(f, delimiter=";"))
        rutas = {r["archivo_origen"]: r["ruta_original"] for r in filas}
        self.assertEqual(rutas["Vigilante.docx"], "entrada/Descriptivos/Seguridad/Vigilante.docx")

    def test_construccion_sin_problemas(self):
        self.assertEqual(self.codigo, 0, self.salida)


if __name__ == "__main__":
    unittest.main()
