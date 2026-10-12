"""Corpus sintetico de prueba (Word + PDF) con trampas de falsos positivos conocidas.

Requiere python-docx y reportlab (solo para las pruebas). Uso directo:
    python tests/fixtures_corpus.py <carpeta>
"""
import sys
from pathlib import Path

import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def desc(path, puesto, reporta, proposito, funciones, header):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_heading("Datos generales", 1)
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Nombre del puesto", puesto), ("Reporta a", reporta), ("Supervisa a", "Analista de Compras")]):
        t.cell(i, 0).text = a; t.cell(i, 1).text = b
    d.add_heading("Propósito del puesto", 1)
    d.add_paragraph(proposito)
    d.add_heading("Funciones", 1)
    for f in funciones:
        d.add_paragraph(f, style="List Bullet")
    d.add_heading("Control de cambios", 1)
    d.add_paragraph("Versión 1 inicial del documento.")
    d.save(path)

def proc(path, header, objetivo, pasos, extra_alcance):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_heading("Objetivo", 1)
    d.add_paragraph(objetivo)
    d.add_heading("Alcance", 1)
    d.add_paragraph("Aplica a todas las compras nacionales.")
    for x in extra_alcance:
        d.add_paragraph(x, style="List Bullet")
    d.add_heading("Responsabilidades", 1)
    d.add_paragraph("El Gerente de Compras aprueba las órdenes. El Jefe de Almacén recibe el material.")
    d.add_heading("Desarrollo del proceso", 1)
    d.add_heading("Solicitud", 2)
    t = d.add_table(rows=len(pasos) + 1, cols=3)
    for j, h in enumerate(["No.", "Actividad", "Responsable"]):
        t.cell(0, j).text = h
    for i, (a, r) in enumerate(pasos, 1):
        t.cell(i, 0).text = str(i); t.cell(i, 1).text = a; t.cell(i, 2).text = r
    d.add_paragraph("Registrar la solicitud en el sistema.", style="List Bullet")
    d.add_paragraph("Confirmar con Almacén la disponibilidad.", style="List Bullet")
    d.add_heading("Documentos de referencia", 1)
    d.add_paragraph("PR-ALM-001 Recepción de materiales; norma ISO-9001; protocolo COVID-19.")
    d.save(path)

def generar(base):
    """Crea <base>/entrada (lista.txt, Descriptivos/, Procedimientos/) y <base>/contexto."""
    base = Path(base)
    ent = base / "entrada"
    ctx = base / "contexto"
    for p in (ent, ctx):
        p.mkdir(parents=True, exist_ok=True)
    compras = ent / "Descriptivos" / "Compras"
    compras.mkdir(parents=True, exist_ok=True)
    desc(compras / "Gerente de Compras.docx", "Gerente de Compras", "Director General",
         "Asegurar el abastecimiento oportuno de materiales.", ["Aprobar órdenes de compra.", "Coordinar con Almacén las entregas.", "Revisar el estado de cuenta de proveedores."],
         "Código: DP-COM-01 Versión: 2")
    desc(compras / "Analista de Compras.docx", "Analista de Compras", "Gerente de Compras",
         "Gestionar cotizaciones. Responsable de mantener el registro de proveedores.", ["Solicitar cotizaciones.", "Enviar órdenes a proveedores."],
         "Código: DP-COM-02 Estado: Vigente")
    alm = ent / "Descriptivos" / "Almacen"
    alm.mkdir(parents=True, exist_ok=True)
    desc(alm / "Jefe de Almacen.docx", "Jefe de Almacén", "Director de Operaciones",
         "Controlar el inventario físico. Estado de cuenta revisado mensualmente.", ["Recibir materiales de Compras.", "Notificar a Compras faltantes."],
         "Código: DP-ALM-01")
    pc = ent / "Procedimientos" / "Compras"
    pc.mkdir(parents=True, exist_ok=True)
    proc(pc / "Compra de materiales.docx", "Código: PR-COM-001 Versión: 3 Estado: Vigente Propietario: Gerente de Compras",
         "Establecer el proceso de compra.", [("Solicitar material", "Jefe de Almacén"), ("Aprobar la orden", "Gerente de Compras"), ("Confirmar la recepción", "Analista de Compras")],
         ["Compras nacionales", "Compras internacionales"])
    proc(pc / "Compra urgente.docx", "Código: PR-COM-002 Estado: Obsoleto",
         "Atender compras urgentes. Revisión de la solicitud por el área.", [("Solicitar compra urgente", "Jefe de Almacén")], [])
    pa = ent / "Procedimientos" / "Almacen"
    pa.mkdir(parents=True, exist_ok=True)

    # PDF de almacen
    c = canvas.Canvas(str(pa / "Recepcion de materiales.pdf"), pagesize=letter)
    for pg in range(2):
        c.setFont("Helvetica", 9)
        c.drawString(72, 750, "Código: PR-ALM-001   Versión: 1   Página %d de 2" % (pg + 1))
        c.line(72, 742, 540, 742)
        y = 700
        if pg == 0:
            c.setFont("Helvetica-Bold", 13); c.drawString(72, y, "1. Objetivo"); y -= 20
            c.setFont("Helvetica", 10); c.drawString(72, y, "Recibir los materiales comprados por Compras y verificar su calidad."); y -= 30
            c.setFont("Helvetica-Bold", 13); c.drawString(72, y, "2. Alcance"); y -= 20
            c.setFont("Helvetica", 10); c.drawString(72, y, "Aplica al almacén central."); y -= 30
            c.setFont("Helvetica-Bold", 13); c.drawString(72, y, "3. Desarrollo del proceso"); y -= 20
            c.setFont("Helvetica", 10)
            for t in ["• El Jefe de Almacén recibe el material.", "• El Gerente de Compras autoriza el pago.", "• Se confirma la entrada al sistema."]:
                c.drawString(80, y, t); y -= 16
        else:
            c.setFont("Helvetica-Bold", 13); c.drawString(72, y, "4. Documentos de referencia"); y -= 20
            c.setFont("Helvetica", 10); c.drawString(72, y, "PR-COM-001 Compra de materiales."); y -= 20
        c.setFont("Helvetica", 8); c.drawString(72, 40, "Documento controlado - uso interno")
        c.showPage()
    c.save()

    lista = []
    for p in sorted(ent.rglob("*")):
        if p.suffix in (".docx", ".pdf"):
            lista.append("C:\\Docs\\" + "\\".join(p.relative_to(ent).parts))
    (ent / "lista.txt").write_text("\n".join(lista) + "\n", encoding="utf-8")

    g = docx.Document()
    g.add_heading("Glosario", 1)
    t = g.add_table(rows=4, cols=2)
    for i, (a, b) in enumerate([("Sigla", "Significado"), ("ALM", "Almacen"), ("Responsable", "Compras"), ("Área", "Compras")]):
        t.cell(i, 0).text = a; t.cell(i, 1).text = b
    g.save(ctx / "Glosario.docx")
    return ent, ctx


if __name__ == "__main__":
    generar(sys.argv[1])
