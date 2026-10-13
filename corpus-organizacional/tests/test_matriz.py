"""Matriz organizacional en Excel."""

import tempfile
import unittest
from pathlib import Path

from tests import ayuda  # noqa: F401
from corpus_org.comun.archivos import escribir_csv
from corpus_org.comun.esquema import COLUMNAS_INDICE
from corpus_org.consulta.matriz import datos_matriz, escribir_matriz, guardar_libro

DOCS = """## [D-PRU-01] Coordinador de Pruebas
> - Código: D-PRU-01

### [D-PRU-01] 1. Coordinador de Pruebas.

Coordinar las pruebas según la Política de Calidad y el procedimiento P-PRU-01.

### [D-PRU-01] Datos de encabezado y pie del documento original

Código: DP-FOR-001

## [D-PRU-02] Analista de Pruebas
> - Código: D-PRU-02

### [D-PRU-02] 1. Analista de Pruebas.

| FUNCIONES ESPECÍFICAS | EVIDENCIA DE CUMPLIMENTO / ENTREGABLES |
|---|---|
| Ejecutar los casos de prueba. | • Casos ejecutados. |
| Registrar hallazgos según el protocolo de pruebas. | • Registro. |

### [D-PRU-02] Áreas de Relacionamiento:

| Área de Interacción | Tipo de Relación |
|---|---|
| Coordinador de Pruebas | Subordinación directa |
| Gerentes de Producto | Equipo Interdisciplinario |

| Requisitos del Puesto |   |
|---|---|
| Nivel de Educación Formal | Ingeniería o afines. |
"""
PROC = """## [P-PRU-01] Ejecución de pruebas
> - Código: P-PRU-01

### [P-PRU-01] Desarrollo del proceso

| No. | Actividad | Responsable |
|---|---|---|
| 1 | Ejecutar los casos de prueba | Analista de Pruebas |
| 2 | Aprobar el informe de pruebas | Coordinador de Pruebas |
"""


def _corpus(raiz):
    pub = raiz / "publicado"
    pub.mkdir(parents=True)
    (pub / "D_PRU.md").write_text(DOCS, encoding="utf-8")
    (pub / "P_PRU.md").write_text(PROC, encoding="utf-8")
    filas = [{"codigo": "D-PRU-01", "tipo": "Descriptivo", "departamento": "Pruebas",
              "nombre_puesto": "Coordinador de Pruebas", "archivo_md": "D_PRU.md"},
             {"codigo": "D-PRU-02", "tipo": "Descriptivo", "departamento": "Pruebas",
              "nombre_puesto": "Analista de Pruebas", "reporta_a": "Coordinador de Pruebas",
              "detalle_avisos": "OCR en páginas 1: revisar contra el original", "archivo_md": "D_PRU.md"},
             {"codigo": "P-PRU-01", "tipo": "Procedimiento", "departamento": "Pruebas",
              "nombre": "Ejecución de pruebas", "archivo_md": "P_PRU.md"}]
    escribir_csv(pub / "indice_corpus.csv", COLUMNAS_INDICE, filas)
    (raiz / "RESUMEN_CONSTRUCCION.md").write_text("- Resultado: **LISTO PARA PUBLICAR**\n", encoding="utf-8")


class MatrizTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.raiz = Path(cls.tmp.name)
        _corpus(cls.raiz)
        _, cls.hojas = datos_matriz(cls.raiz)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_puestos_y_jefe_del_corpus(self):
        analista = next(f for f in self.hojas["Puestos"] if f[0] == "D-PRU-02")
        self.assertEqual((analista[3], analista[4]), ("Coordinador de Pruebas", "D-PRU-01"))
        self.assertEqual(analista[12], "OCR (confianza desconocida; revisar contra el original)")

    def test_funciones_relaciones_y_perfil(self):
        self.assertEqual([f[4:] for f in self.hojas["Funciones"]],
                         [["Ejecutar los casos de prueba.", "• Casos ejecutados."],
                          ["Registrar hallazgos según el protocolo de pruebas.", "• Registro."]])
        self.assertIn(["D-PRU-02", "Pruebas", "Analista de Pruebas", "Coordinador de Pruebas",
                       "Subordinación directa", "D-PRU-01", "Coordinador de Pruebas"], self.hojas["Relaciones"])
        self.assertIn(["D-PRU-02", "Pruebas", "Analista de Pruebas", "Requisitos del Puesto",
                       "Nivel de Educación Formal", "Ingeniería o afines."], self.hojas["Perfil y requisitos"])

    def test_normativa_citada(self):
        tipos = {(f[0], f[3]) for f in self.hojas["Normativa citada"]}
        self.assertTrue({("D-PRU-01", "Política"), ("D-PRU-01", "Procedimiento"), ("D-PRU-02", "Protocolo")} <= tipos)
        # el codigo del encabezado del formato no es una cita
        self.assertFalse(any("DP-FOR-001" in f[4] for f in self.hojas["Normativa citada"]))

    def test_participacion_en_procedimientos(self):
        roles = {(f[0], f[2]) for f in self.hojas["Participación"]}
        self.assertEqual(roles, {("D-PRU-02", "P-PRU-01"), ("D-PRU-01", "P-PRU-01")})
        matriz = self.hojas["Puestos x Procedimientos"]
        self.assertEqual(matriz[0], ["Código", "Puesto", "P-PRU-01"])

    def test_libro(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("falta openpyxl")
        ruta = escribir_matriz(self.raiz)
        wb = openpyxl.load_workbook(ruta)
        self.assertEqual(wb.sheetnames[:3], ["Léeme", "Resumen", "Puestos"])
        self.assertEqual(wb["Resumen"]["B2"].value, "=COUNTIF('Puestos'!$B:$B,$A2)")
        self.assertEqual(wb["Funciones"].max_row, 3)


class GuardarLibroTest(unittest.TestCase):
    def test_archivo_abierto_en_excel(self):
        """Si el Excel anterior esta abierto, la matriz nueva se guarda al lado con fecha y hora."""
        guardados = []

        class Libro:
            def save(self, ruta):
                if ruta.name == "matriz_organizacional.xlsx":
                    raise PermissionError(13, "Permission denied")
                guardados.append(ruta)

        destino = Path(tempfile.gettempdir()) / "matriz_organizacional.xlsx"
        ruta = guardar_libro(Libro(), destino)
        self.assertEqual(guardados, [ruta])
        self.assertRegex(ruta.name, r"^matriz_organizacional \d{4}-\d{2}-\d{2} \d{6}\.xlsx$")


if __name__ == "__main__":
    unittest.main()
