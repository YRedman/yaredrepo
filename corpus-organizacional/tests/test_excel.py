"""Lectura de Excel (.xlsx) y deteccion del formato real por contenido.

Los libros se generan con datos ficticios; ningun documento real entra al repositorio.
"""

import datetime
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from tests import ayuda  # noqa: F401

try:
    import openpyxl
    from openpyxl.drawing.image import Image as ImagenXl
    from PIL import Image, ImageDraw, ImageFont
    from corpus_org.conversion.extractores import extractor_para, formato_real
    from corpus_org.conversion.extractores.excel import valor_texto
    from corpus_org.conversion.extractores.ocr import idioma_ocr
    DEPENDENCIAS = True
except ImportError:
    DEPENDENCIAS = False

OCR = DEPENDENCIAS and idioma_ocr() is not None
FUENTE = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def _libro(ruta, construir):
    wb = openpyxl.Workbook()
    construir(wb)
    wb.save(ruta)


def _extraer(ruta):
    return extractor_para(ruta).extraer(ruta)


@unittest.skipUnless(DEPENDENCIAS, "faltan openpyxl o Pillow")
class ExcelTest(unittest.TestCase):
    def test_una_hoja_con_fechas_porcentajes_y_enteros_es_una_tabla(self):
        def construir(wb):
            ws = wb.active
            ws.title = "Feriados"
            ws.append(["Fecha", "Cumplimiento", "Horas", "Nota"])
            ws.append([datetime.datetime(2026, 1, 1), 0.85, 8.0, "Año Nuevo"])
            ws.append([datetime.date(2026, 1, 6), 0.5, 7.5, ""])
            ws["B2"].number_format = "0%"
            ws["B3"].number_format = "0%"
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "feriados.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        tablas = [b["filas"] for b in d.bloques if b["k"] == "tabla"]
        self.assertEqual(tablas, [[["Fecha", "Cumplimiento", "Horas", "Nota"], ["2026-01-01", "85%", "8", "Año Nuevo"],
                                   ["2026-01-06", "50%", "7.5", ""]]])
        self.assertEqual(d.tablas, 1)
        self.assertFalse(d.avisos)

    def test_varias_hojas_llevan_su_titulo_y_las_ocultas_se_avisan(self):
        def construir(wb):
            a = wb.active
            a.title = "Ventas"
            a.append(["Mes", "Total"])
            a.append(["Enero", 10])
            b = wb.create_sheet("Gastos")
            b.append(["Mes", "Total"])
            b.append(["Enero", 4])
            o = wb.create_sheet("Secreta")
            o.append(["x", "y"])
            o.sheet_state = "hidden"
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "libro.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        titulos = [b["texto"] for b in d.bloques if b["k"] == "h"]
        self.assertEqual(titulos, ["Ventas", "Gastos"])
        self.assertTrue(any("Secreta" in n for n in d.notas), d.notas)

    def test_una_hoja_de_una_columna_son_parrafos(self):
        def construir(wb):
            ws = wb.active
            ws.append(["Primer renglon del texto"])
            ws.append(["Segundo renglon del texto"])
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "texto.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        self.assertEqual([b["texto"] for b in d.bloques if b["k"] == "p"],
                         ["Primer renglon del texto", "Segundo renglon del texto"])

    def test_una_hoja_con_varias_zonas_queda_en_tablas_y_titulos_separados(self):
        def construir(wb):
            ws = wb.active
            ws.title = "Ficha"
            ws["B2"], ws["C2"] = "Proceso:", "Compras"
            ws["B3"], ws["C3"] = "Lider:", "Director"
            ws["B5"], ws["D5"], ws["F5"] = "Proveedor", "Actividad", "Salida"       # tabla con columnas vacias entre medio
            ws["B6"], ws["D6"], ws["F6"] = "Almacen", "Recibir", "Orden"
            ws["D7"], ws["F7"] = "Registrar", "Registro"
            ws["B10"] = "RECURSOS"
            ws["B11"], ws["D11"] = "Actividad", "Recurso"
            ws["B12"], ws["D12"] = "Recibir", "Escaner"
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "ficha.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        self.assertEqual([b["texto"] for b in d.bloques if b["k"] == "h"], ["RECURSOS"])
        tablas = [b["filas"] for b in d.bloques if b["k"] == "tabla"]
        self.assertEqual(tablas, [[["Proceso:", "Compras"], ["Lider:", "Director"]],
                                  [["Proveedor", "Actividad", "Salida"], ["Almacen", "Recibir", "Orden"],
                                   ["", "Registrar", "Registro"]],
                                  [["Actividad", "Recurso"], ["Recibir", "Escaner"]]])
        self.assertEqual(d.tablas, 3)

    def test_una_fila_en_blanco_dentro_de_la_misma_tabla_no_la_parte(self):
        def construir(wb):
            ws = wb.active
            ws.append(["Id", "Nombre"])
            ws.append([1, "Uno"])
            ws.append([])
            ws.append([2, "Dos"])
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "t.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        self.assertEqual([b["filas"] for b in d.bloques if b["k"] == "tabla"],
                         [[["Id", "Nombre"], ["1", "Uno"], ["2", "Dos"]]])

    def test_celdas_combinadas_en_vertical_repiten_su_valor_pero_un_rotulo_ancho_no(self):
        def construir(wb):
            ws = wb.active
            ws["B1"] = "TITULO DEL FORMULARIO"
            ws.merge_cells("B1:H2")
            ws.append([])
            ws["A4"], ws["B4"] = "Area", "Tarea"
            ws["A5"], ws["B5"] = "Ventas", "Cotizar"
            ws["B6"] = "Cobrar"
            ws.merge_cells("A5:A6")
            ws["A7"], ws["B7"] = "Compras", "Pedir"
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "m.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        self.assertEqual([b["texto"] for b in d.bloques if b["k"] == "p"], ["TITULO DEL FORMULARIO"])
        self.assertEqual([b["filas"] for b in d.bloques if b["k"] == "tabla"],
                         [[["Area", "Tarea"], ["Ventas", "Cotizar"], ["Ventas", "Cobrar"], ["Compras", "Pedir"]]])

    def test_formulas_sin_valor_guardado_se_avisan(self):
        def construir(wb):
            ws = wb.active
            ws.append(["a", "b", "suma"])
            ws.append([1, 2, "=A2+B2"])      # openpyxl no guarda el valor calculado
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "formulas.xlsx"
            _libro(ruta, construir)
            d = _extraer(ruta)
        self.assertTrue(any("fórmulas sin valor guardado" in a for a in d.avisos), d.avisos)
        self.assertEqual([b["filas"] for b in d.bloques if b["k"] == "tabla"],
                         [[["a", "b", "suma"], ["1", "2", "=A2+B2"]]])      # la formula en vez de una celda vacia
        self.assertTrue(d.hoja_de_calculo)

    def test_un_excel_con_extension_csv_se_lee_como_excel(self):
        def construir(wb):
            ws = wb.active
            ws.append(["Mes", "Total"])
            ws.append(["Enero", 10])
        with tempfile.TemporaryDirectory() as tmp:
            real = Path(tmp) / "datos.xlsx"
            _libro(real, construir)
            falso = Path(tmp) / "datos.csv"
            falso.write_bytes(real.read_bytes())
            self.assertEqual(formato_real(falso), ".xlsx")
            d = _extraer(falso)
        self.assertEqual([b["filas"] for b in d.bloques if b["k"] == "tabla"], [[["Mes", "Total"], ["Enero", "10"]]])
        self.assertTrue(any("es un libro de Excel" in n for n in d.notas), d.notas)

    def test_valor_texto(self):
        ws = openpyxl.Workbook().active
        ws["A1"] = 3.0
        ws["A2"] = 0.125
        ws["A3"] = True
        ws["A4"] = "  dos   palabras "
        self.assertEqual([valor_texto(ws[c]) for c in ("A1", "A2", "A3", "A4", "A5")],
                         ["3", "0.125", "VERDADERO", "dos palabras", ""])


@unittest.skipUnless(DEPENDENCIAS, "faltan openpyxl o Pillow")
class FormatoRealTest(unittest.TestCase):
    def test_se_reconoce_por_el_contenido(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "a.txt").write_bytes(b"%PDF-1.4\n...")
            (t / "b.csv").write_bytes(b"a;b\n1;2\n")
            (t / "c.csv").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as z:
                z.writestr("ppt/presentation.xml", "<p/>")
            (t / "d.csv").write_bytes(buf.getvalue())
            self.assertEqual(formato_real(t / "a.txt"), ".pdf")
            self.assertIsNone(formato_real(t / "b.csv"))
            self.assertEqual(formato_real(t / "c.csv"), ".ole")
            self.assertEqual(formato_real(t / "d.csv"), ".pptx")
            with self.assertRaises(RuntimeError) as e:
                extractor_para(t / "c.csv")
            self.assertIn("Office antiguo", str(e.exception))


@unittest.skipUnless(DEPENDENCIAS, "faltan openpyxl o Pillow")
class AvisosDeHojaTest(unittest.TestCase):
    def test_celda_inicial_vacia_es_un_dato_en_una_hoja_y_una_fila_partida_en_un_pdf(self):
        from corpus_org.conversion.analisis.estructura import avisos_de_estructura
        filas = [["Área", "Meta"]] + [[f"Área {k}", str(k)] for k in range(1, 8)] + [["", "20"]]
        tabla = {"k": "tabla", "filas": filas}
        self.assertTrue(any("sin primera celda" in a for a in avisos_de_estructura([tabla], 0)))
        self.assertEqual(avisos_de_estructura([tabla], 0, hoja_de_calculo=True), [])

    def test_ruta_original_sale_del_archivo_realmente_leido(self):
        from types import SimpleNamespace
        from corpus_org.conversion.lote import ruta_original
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "contexto" / "Plan").mkdir(parents=True)
            (t / "entrada").mkdir()
            archivo = t / "contexto" / "Plan" / "Procedimiento.pdf"
            archivo.write_bytes(b"x")
            opciones = SimpleNamespace(raiz_entrada=t / "entrada", raiz_contexto=t / "contexto")
            fila = {"ruta_original": "corpus-organizacional/contexto/Plan/Procedimiento.pdf", "archivo": "Procedimiento.pdf"}
            grupo = SimpleNamespace(es_contexto=False)
            self.assertEqual(ruta_original(fila, grupo, opciones, archivo), "contexto/Plan/Procedimiento.pdf")
            # sin el archivo (no adjunto) se usa la ruta del inventario
            self.assertEqual(ruta_original(fila, grupo, opciones), "entrada/contexto/Plan/Procedimiento.pdf")


@unittest.skipUnless(OCR, "faltan Tesseract (spa), pytesseract o Pillow")
class ImagenEnExcelTest(unittest.TestCase):
    def test_el_texto_de_una_captura_pegada_en_la_hoja_se_lee_con_ocr(self):
        img = Image.new("RGB", (640, 420), "white")
        dib = ImageDraw.Draw(img)
        try:
            f = ImageFont.truetype(FUENTE, 20)
        except OSError:
            f = ImageFont.load_default()
        for k, t in enumerate(["Jueves primero de enero dia de Año Nuevo", "Lunes veintiseis de enero dia de Duarte",
                               "Viernes veintisiete de febrero Independencia", "Viernes tres de abril Viernes Santo",
                               "Domingo dieciseis de agosto Restauracion"]):
            dib.text((20, 30 + 50 * k), t, fill="black", font=f)
        with tempfile.TemporaryDirectory() as tmp:
            png = Path(tmp) / "captura.png"
            img.save(png)
            ruta = Path(tmp) / "con_imagen.xlsx"

            def construir(wb):
                wb.active.title = "Calendario"
                wb.active.add_image(ImagenXl(str(png)), "B2")
            _libro(ruta, construir)
            d = _extraer(ruta)
        texto = " ".join(b.get("texto", "") for b in d.bloques)
        self.assertIn("Duarte", texto)
        self.assertIn("texto leído con OCR", texto)
        self.assertEqual(d.paginas_ocr, [1])
        self.assertGreater(d.confianza_ocr, 70)


if __name__ == "__main__":
    unittest.main()
