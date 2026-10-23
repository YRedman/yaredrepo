"""Extractor de libros de Excel (.xlsx, .xlsm) usados como material de contexto.

Cada hoja visible se divide en bloques por las filas en blanco: cada bloque es una tabla (su primera fila es la
cabecera), un titulo de seccion o, si tiene una sola columna, parrafos. Las celdas combinadas en vertical repiten su
valor en cada fila. Se leen los valores guardados, no las formulas. Las imagenes pegadas en una hoja (una captura de un
documento, un organigrama) se leen con OCR y quedan bajo una marca, igual que en un PDF.
"""

import collections
import datetime
import importlib.util
import io
import warnings
import zipfile
from pathlib import Path

from ...comun.esquema import RE_MARCAS
from ...comun.texto import numeros, tokens
from ..modelo import DocumentoExtraido, bloque_h
from .base import Extractor
from .imagenes import es_imagen_de_contenido, texto_de_imagen

MAX_IMAGENES_LIBRO = 20      # imagenes del libro cuyo texto se lee con OCR
FILAS_HOJA_GRANDE = 3000     # a partir de aqui se avisa en las notas que la hoja es grande


def valor_texto(celda):
    """Texto de una celda tal como se ve: fechas ISO, enteros sin ``.0``, porcentajes con %."""
    v = celda.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)):
        if "%" in (celda.number_format or ""):
            return f"{v * 100:.6g}%"
        if isinstance(v, float):
            if v.is_integer():
                return str(int(v))
            return f"{v:.8f}".rstrip("0").rstrip(".")
        return str(v)
    return " ".join(str(v).split())


def _texto_de_celda(c, formulas):
    t = valor_texto(c)
    if not t and formulas:
        t = formulas.get((c.row, c.column), "")     # formula sin valor guardado: mejor la formula que una celda vacia
    return t


MAX_ALTO_COMBINADA = 50      # una combinada vertical mas alta que esto no se rellena hacia abajo
MAX_ANCHO_COMBINADA = 3      # una combinada mas ancha es un rotulo de pagina (titulo del formulario): no se repite
JACCARD_MISMA_TABLA = 0.8    # parecido de columnas usadas para unir dos bloques separados por UNA fila en blanco


def _grilla_de_hoja(ws, formulas=None):
    """Matriz de textos de la hoja. Las celdas combinadas en vertical repiten su valor en cada fila de la combinacion
    (solo en la primera columna), para que cada fila se entienda sola al recuperarla."""
    grilla = [[_texto_de_celda(c, formulas) for c in fila] for fila in ws.iter_rows()]
    for rango in ws.merged_cells.ranges:
        alto = rango.max_row - rango.min_row + 1
        if alto < 2 or alto > MAX_ALTO_COMBINADA or rango.max_col - rango.min_col + 1 > MAX_ANCHO_COMBINADA:
            continue
        f0, c0 = rango.min_row - 1, rango.min_col - 1
        if f0 >= len(grilla) or c0 >= len(grilla[f0]) or not grilla[f0][c0]:
            continue
        for f in range(f0 + 1, min(rango.max_row, len(grilla))):
            if c0 < len(grilla[f]) and not grilla[f][c0]:
                grilla[f][c0] = grilla[f0][c0]
    return grilla


def _columnas(filas):
    return {j for f in filas for j, c in enumerate(f) if c}


def _es_fila_titulo(fila, ancho):
    """Fila con una sola celda llena en un bloque de varias columnas: es un titulo de seccion, no una cabecera."""
    return ancho > 1 and sum(1 for c in fila if c) == 1


def _bloques_de_grilla(grilla):
    """Divide la matriz en bloques de filas separados por filas en blanco. Dos bloques separados por una sola fila en
    blanco con las mismas columnas son la misma tabla. Devuelve [(fila_inicial, [filas])]."""
    grupos, actual = [], None
    for i, f in enumerate(grilla):
        if any(f):
            if actual is None:
                actual = (i, [f])
                grupos.append(actual)
            else:
                actual[1].append(f)
        else:
            actual = None
    bloques = []
    for ini, filas in grupos:
        if bloques:
            ini_p, filas_p = bloques[-1]
            a, b = _columnas(filas_p), _columnas(filas)
            mismo = len(a & b) / max(len(a | b), 1) >= JACCARD_MISMA_TABLA
            if ini - (ini_p + len(filas_p)) == 1 and mismo and not _es_fila_titulo(filas[0], len(b)):
                filas_p.extend(filas)
                continue
        bloques.append((ini, filas))
    return bloques


def _tabla_de_bloque(filas):
    """Filas del bloque sin columnas vacias y con el mismo ancho."""
    ancho = max(len(f) for f in filas)
    filas = [f + [""] * (ancho - len(f)) for f in filas]
    usadas = sorted(_columnas(filas))
    return [[f[j] for j in usadas] for f in filas]


def bloques_de_hoja(ws, formulas=None):
    """Bloques de documento de una hoja: tablas, parrafos (bloque de una columna) y titulos de seccion (fila con una
    sola celda al inicio de un bloque de varias columnas). Una hoja con varias zonas (cabecera del formulario, tabla
    principal, cuadros auxiliares) queda en tablas separadas en vez de una sola rejilla."""
    salida = []
    for _ini, filas in _bloques_de_grilla(_grilla_de_hoja(ws, formulas)):
        tabla = _tabla_de_bloque(filas)
        ancho = len(tabla[0])
        while len(tabla) > 1 and _es_fila_titulo(tabla[0], ancho):
            texto = next(c for c in tabla[0] if c)
            salida.append(bloque_h(2, texto) if any(ch.isalpha() for ch in texto) else {"k": "p", "texto": texto})
            tabla = tabla[1:]
        if ancho == 1 or (len(tabla) == 1 and sum(1 for c in tabla[0] if c) == 1):
            salida += [{"k": "p", "texto": next(c for c in f if c)} for f in tabla]
        else:
            salida.append({"k": "tabla", "filas": tabla})
    return salida


def abrir_libro(path, data_only=True):
    """Libro de openpyxl desde los bytes del archivo: openpyxl se niega a abrir por la extension (un Excel guardado
    como .csv), pero lee igual desde un flujo."""
    import openpyxl
    with warnings.catch_warnings():      # "Conditional Formatting extension is not supported": no afecta a los valores
        warnings.simplefilter("ignore", UserWarning)
        return openpyxl.load_workbook(io.BytesIO(Path(path).read_bytes()), data_only=data_only)


def _textos_de_referencia(wb):
    """Textos de las celdas del libro (4 a 150 caracteres): sirven para corregir lo que el OCR lee en sus diagramas."""
    textos = set()
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and 4 <= len(c.value.strip()) <= 150:
                    textos.add(" ".join(c.value.split()))
    return sorted(textos)


def _imagenes_de_hoja(ws):
    """[bytes] de las imagenes que hay en la hoja (openpyxl las conserva con Pillow instalado)."""
    res = []
    for im in getattr(ws, "_images", []):
        try:
            res.append(im._data())
        except Exception:
            continue
    return res


class ExtractorExcel(Extractor):
    extensiones = (".xlsx", ".xlsm")

    def extraer(self, path):
        if importlib.util.find_spec("openpyxl") is None:
            raise RuntimeError("Falta openpyxl para leer Excel: python -m pip install openpyxl")
        path = Path(path)
        try:
            wb = abrir_libro(path)
        except zipfile.BadZipFile:
            raise RuntimeError("El archivo no es un libro de Excel válido (¿protegido con contraseña o dañado?).")
        d = DocumentoExtraido()
        d.unidad = "hojas"
        d.hoja_de_calculo = True
        textos, bloques = [], []
        ocultas = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
        visibles = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        d.paginas = len(visibles)
        imagenes_leidas = 0
        referencia = None
        sin_valor, formulas = self._formulas_sin_valor(path, visibles, wb)
        for n, ws in enumerate(visibles, 1):
            hoja = bloques_de_hoja(ws, formulas.get(ws.title))
            imagenes = _imagenes_de_hoja(ws)
            if not hoja and not imagenes:
                continue
            if len(visibles) > 1 or imagenes:
                bloques.append(bloque_h(1, ws.title))
                textos.append(ws.title)
            for b in hoja:
                bloques.append(b)
                if b["k"] == "tabla":
                    d.tablas += 1
                    textos += [c for f in b["filas"] for c in f]
                    if len(b["filas"]) > FILAS_HOJA_GRANDE:
                        d.notas.append(f"La hoja «{ws.title}» tiene {len(b['filas'])} filas: considere resumirla.")
                else:
                    textos.append(b["texto"])
            for datos in imagenes:
                if not es_imagen_de_contenido(datos):
                    continue
                if imagenes_leidas >= MAX_IMAGENES_LIBRO:
                    d.imagenes_omitidas += 1
                    continue
                imagenes_leidas += 1
                if referencia is None:
                    referencia = _textos_de_referencia(wb)
                texto_de_imagen(d, bloques, f"LA HOJA «{ws.title}»", n, datos, referencia)
            charts = len(getattr(ws, "_charts", []))
            if charts:
                d.notas.append(f"La hoja «{ws.title}» tiene {charts} gráfico(s): no se leen.")
                bloques.append({"k": "marca", "texto": f"[DIAGRAMA O IMAGEN: gráfico en la hoja «{ws.title}»; "
                                                        "ver el Excel original]"})
        if ocultas:
            d.notas.append("Hojas ocultas que no se leyeron: " + ", ".join(ocultas) + ".")
        if sin_valor:
            d.avisos.append(f"{sin_valor} fórmulas sin valor guardado (se muestran como fórmula, no como resultado): "
                            "abra el archivo en Excel, guárdelo y vuelva a convertir")
        d.bloques = bloques
        crudo = " ".join(textos)
        d.crudo = collections.Counter(tokens(RE_MARCAS.sub(" ", crudo)))
        d.numeros = numeros(crudo)
        if path.suffix.lower() not in self.extensiones:
            d.notas.append(f"El archivo se llama «{path.name}» pero es un libro de Excel: se leyó como Excel.")
        return d

    @staticmethod
    def _formulas_sin_valor(path, visibles, wv):
        """(cantidad, {hoja: {(fila, columna): formula}}) de las formulas de las hojas visibles que no tienen valor
        guardado (se leerian como celdas vacias)."""
        try:
            wf = abrir_libro(path, data_only=False)
        except Exception:
            return 0, {}
        total, por_hoja = 0, {}
        for ws in visibles:
            if ws.title not in wf.sheetnames:
                continue
            for fila_f, fila_v in zip(wf[ws.title].iter_rows(), wv[ws.title].iter_rows()):
                for cf, cv in zip(fila_f, fila_v):
                    if isinstance(cf.value, str) and cf.value.startswith("=") and cv.value is None:
                        total += 1
                        por_hoja.setdefault(ws.title, {})[(cf.row, cf.column)] = cf.value
        return total, por_hoja
