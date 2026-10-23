"""El libro de Excel de la matriz (requiere openpyxl)."""

import datetime

from ... import VERSION
from ...comun.errores import CorpusError
from .esquema import COLUMNAS_RESUMEN, DESCRIPCIONES, cabecera_y_filas
from .exportacion import estado_construccion
from .resumen import COLUMNAS_TOTAL, areas_del_resumen, formulas_resumen

COLOR_CABECERA, COLOR_AUDITOR = "1F4E78", "FFFF00"
COLUMNA_AUDITOR = 7


def guardar_libro(wb, destino):
    """Guarda el libro; si el archivo esta abierto (Excel lo bloquea en Windows) lo guarda con otro nombre.

    Devuelve la ruta donde quedo guardado. Asi una construccion larga no termina en error por un
    archivo abierto: la matriz nueva queda al lado, con la fecha y hora en el nombre.
    """
    try:
        wb.save(destino)
        return destino
    except PermissionError:
        alterno = destino.with_name(f"{destino.stem} {datetime.datetime.now():%Y-%m-%d %H%M%S}{destino.suffix}")
        wb.save(alterno)
        return alterno


def _importar_openpyxl():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise CorpusError("Falta openpyxl para escribir el Excel: python -m pip install openpyxl")
    return openpyxl, Alignment, Font, PatternFill, get_column_letter


def escribir_excel(hojas, raiz, destino):
    """Escribe el libro con una hoja por cada una de ``hojas``, mas Léeme y Resumen. Devuelve la ruta guardada."""
    openpyxl, Alignment, Font, PatternFill, get_column_letter = _importar_openpyxl()
    fuente, negrita = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    relleno_cab = PatternFill("solid", fgColor=COLOR_CABECERA)
    relleno_auditor = PatternFill("solid", fgColor=COLOR_AUDITOR)
    ajuste = Alignment(wrap_text=True, vertical="top")

    wb = openpyxl.Workbook()
    leeme = wb.active
    leeme.title = "Léeme"
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        cabecera, filas, anchos = cabecera_y_filas(nombre, filas)
        ws.append(cabecera)
        for f in filas:
            ws.append(f)
        for j, a in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(j)].width = a
        for fila in ws.iter_rows():
            for celda in fila:
                celda.font, celda.alignment = fuente, ajuste
        for celda in ws[1]:
            celda.font, celda.fill = negrita, relleno_cab
        if nombre == "Revisión":
            for k in range(2, ws.max_row + 1):
                ws.cell(k, COLUMNA_AUDITOR).fill = relleno_auditor
        ws.freeze_panes = "A2" if nombre != "Puestos x Procedimientos" else "C2"
        ws.auto_filter.ref = ws.dimensions

    _leeme(leeme, raiz, fuente)
    _resumen(wb.create_sheet("Resumen", 1), hojas, fuente)
    wb.calculation.fullCalcOnLoad = True
    return guardar_libro(wb, destino)


def _leeme(ws, raiz, fuente):
    from openpyxl.styles import Alignment, Font
    titulo = Font(name="Arial", size=14, bold=True)
    negrita = Font(name="Arial", size=10, bold=True)
    filas = [
        ("Matriz organizacional", titulo),
        (f"Generada el {datetime.date.today():%d/%m/%Y} con convertir.py v{VERSION} a partir del corpus publicado.",
         fuente),
        (f"Estado de la construcción del corpus: {estado_construccion(raiz)}.", negrita),
        ("Todo el contenido es literal de los documentos. Las columnas «Código del jefe en el corpus» y «Código "
         "relacionado» son cruces automáticos por nombre de puesto; «probable» indica que el nombre no es idéntico.",
         fuente),
        ("«Calidad de lectura» = Digital (texto del PDF), OCR (documento escaneado: revisar contra el original) o "
         "Sin texto. Los documentos «Pendiente» en la hoja Revisión aún no se compararon con el original.", fuente),
        ("", fuente),
        ("Hojas", negrita),
    ]
    for texto, f in filas:
        ws.append([texto])
        ws.cell(ws.max_row, 1).font = f
    for nombre, desc in DESCRIPCIONES:
        ws.append([nombre, desc])
        ws.cell(ws.max_row, 1).font = negrita
        ws.cell(ws.max_row, 2).font = fuente
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    for fila in ws.iter_rows():
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")


def _resumen(ws, hojas, fuente):
    """Conteos por área con fórmulas sobre las otras hojas."""
    from openpyxl.styles import Font, PatternFill
    ws.append(COLUMNAS_RESUMEN)
    for celda in ws[1]:
        celda.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_CABECERA)
    for r, area in enumerate(areas_del_resumen(hojas), 2):
        ws.append([area] + formulas_resumen(r))
    ultima = ws.max_row
    ws.append(["Total"] + [f"=SUM({chr(ord('A') + k)}2:{chr(ord('A') + k)}{ultima})" for k in COLUMNAS_TOTAL])
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font = fuente
    for celda in ws[ws.max_row]:
        celda.font = Font(name="Arial", size=10, bold=True)
    ws.column_dimensions["A"].width = 30
    for k in COLUMNAS_TOTAL:
        ws.column_dimensions[chr(ord("A") + k)].width = 16
    ws.freeze_panes = "A2"
